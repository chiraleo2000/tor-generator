"""Live hire_develop workflow quality check: gold pack → draft → export → score."""

from __future__ import annotations

import json
import sys
import time
from datetime import datetime
from pathlib import Path

import httpx
from docx import Document

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))
from app.export.draft_quality import score_draft_text  # noqa: E402

API = "http://127.0.0.1:4000"
EMAIL = "officer@example.go.th"
PASSWORD = "Passw0rd!"
GOLD = BACKEND / "tests" / "fixtures" / "hire_develop_gold_pack.txt"
OUT_DIR = Path(__file__).resolve().parents[3] / "Discussions" / "test-evidence"


def _step(msg: str) -> None:
    print(f"[{datetime.now():%H:%M:%S}] {msg}", flush=True)


def _data(response: httpx.Response) -> dict:
    body = response.json()
    if isinstance(body, dict) and "data" in body:
        return body["data"] or {}
    return body if isinstance(body, dict) else {}


def _login(client: httpx.Client) -> None:
    response = client.post("/api/v1/auth/login", json={"email": EMAIL, "password": PASSWORD})
    response.raise_for_status()
    token = _data(response).get("token")
    if token:
        client.headers["Authorization"] = f"Bearer {token}"


def _wait_healthy(client: httpx.Client) -> None:
    for _ in range(30):
        try:
            health = client.get("/health", timeout=5.0)
            if health.status_code == 200:
                return
        except httpx.HTTPError:
            time.sleep(2)
    raise SystemExit("API not healthy")


def _docx_text(path: Path) -> str:
    doc = Document(str(path))
    paras = [p.text for p in doc.paragraphs]
    tables = [
        cell.text
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
    ]
    return "\n".join([*paras, *tables])


def _create_project(client: httpx.Client) -> str:
    created = client.post(
        "/api/v1/projects",
        json={
            "name": f"workflow quality {datetime.now():%Y%m%d-%H%M%S}",
            "ministry": "กรมบัญชีกลาง",
            "budget": 2_500_000,
            "project_type": "hire_develop",
        },
    )
    created.raise_for_status()
    return str(_data(created)["id"])


def _intake_and_analyze(client: httpx.Client, project_id: str) -> list[str]:
    pack = GOLD.read_text(encoding="utf-8")
    pasted = client.post(
        f"/api/v1/projects/{project_id}/intake/text",
        json={"content": pack},
    )
    pasted.raise_for_status()
    _step("pasted gold pack")
    analyzed = client.post(f"/api/v1/projects/{project_id}/intake/analyze")
    analyzed.raise_for_status()
    coverage = _data(analyzed).get("coverage") or []
    filled = [
        row.get("key")
        for row in coverage
        if row.get("filled") or row.get("status") == "filled"
    ]
    _step(f"analyze filled {len(filled)}/{len(coverage)} {filled[:16]}")
    refs = client.post(f"/api/v1/projects/{project_id}/intake/fill-references")
    _step(f"fill-references {refs.status_code}")
    refs.raise_for_status()
    ready = client.post(
        f"/api/v1/projects/{project_id}/intake/confirm-ready",
        json={"confirm": True},
    )
    ready.raise_for_status()
    return [str(key) for key in filled if key]


def _kickoff_draft(client: httpx.Client, project_id: str) -> None:
    try:
        with client.stream(
            "POST",
            f"/api/v1/projects/{project_id}/draft-chat/start",
            timeout=httpx.Timeout(45.0, connect=10.0),
        ) as response:
            if response.status_code >= 400:
                raise SystemExit(f"draft-chat/start {response.status_code} {response.read()[:400]!r}")
            next(response.iter_lines(), None)
    except httpx.TimeoutException:
        _step("draft-chat/start stream timed out — job should continue")


def _wait_all_drafted(client: httpx.Client, project_id: str) -> dict:
    deadline = time.time() + 3600
    last: dict = {}
    while time.time() < deadline:
        status = client.get(f"/api/v1/projects/{project_id}/draft-chat/status")
        last = _data(status)
        drafted = last.get("drafted_count")
        total = last.get("total")
        _step(f"draft {drafted}/{total} job={last.get('job_status')} all={last.get('all_drafted')}")
        if last.get("all_drafted") and int(total or 0) > 0:
            return last
        time.sleep(8)
    raise SystemExit(f"draft timed out last={last}")


def _export_docx(client: httpx.Client, project_id: str) -> bytes:
    export = client.post(
        f"/api/v1/projects/{project_id}/export",
        json={"use_thai_numerals": True, "numbering_scheme": "none"},
    )
    _step(f"export start {export.status_code}")
    for _ in range(60):
        status = client.get(f"/api/v1/projects/{project_id}/export/status")
        payload = _data(status)
        _step(f"export status={payload.get('status')}")
        if payload.get("status") == "completed":
            return client.get(f"/api/v1/projects/{project_id}/export/download/docx").content
        if payload.get("status") in {"failed", "error"}:
            break
        time.sleep(3)
    return b""


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if not GOLD.exists():
        raise SystemExit(f"missing gold pack {GOLD}")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    with httpx.Client(base_url=API, timeout=httpx.Timeout(1800.0, connect=30.0)) as client:
        _wait_healthy(client)
        _login(client)
        project_id = _create_project(client)
        _step(f"project {project_id}")
        filled = _intake_and_analyze(client, project_id)
        _kickoff_draft(client, project_id)
        status = _wait_all_drafted(client, project_id)
        confirm4 = client.post(
            f"/api/v1/projects/{project_id}/intake/confirm-phase4",
            json={"confirm": True},
        )
        _step(f"confirm-phase4 {confirm4.status_code}")
        confirm4.raise_for_status()
        docx_bytes = _export_docx(client, project_id)
        out_docx = OUT_DIR / "workflow-quality-export.docx"
        if not docx_bytes.startswith(b"PK"):
            report = {
                "project_id": project_id,
                "error": "no docx",
                "draft_status": status,
                "filled": filled,
            }
            (OUT_DIR / "_round-workflow-quality.json").write_text(
                json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
            )
            return 1
        out_docx.write_bytes(docx_bytes)
        text = _docx_text(out_docx)
        harness = score_draft_text(text, category="hire_develop")
        report = {
            "project_id": project_id,
            "filled": filled,
            "draft_status": {
                "drafted_count": status.get("drafted_count"),
                "total": status.get("total"),
                "all_drafted": status.get("all_drafted"),
            },
            "harness": harness,
            "docx": str(out_docx),
        }
        dest = OUT_DIR / "_round-workflow-quality.json"
        dest.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        _step(f"wrote {dest}")
        print(json.dumps(harness, ensure_ascii=False, indent=2))
        bad = bool(harness.get("filename_leak") or harness.get("type_forbidden"))
        if int(status.get("total") or 0) != 16 or not status.get("all_drafted"):
            bad = True
        return 1 if bad else 0


if __name__ == "__main__":
    raise SystemExit(main())
