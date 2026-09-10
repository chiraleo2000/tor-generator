"""Unit + Hypothesis tests for tor-output-standardization (Properties 1–13)."""

from __future__ import annotations

import inspect

import pytest
from hypothesis import given, settings, strategies as st

from app.domain.tor_taxonomy import (
    CRITICAL_SECTIONS_MIN_LENGTH,
    CORE_SECTION_ORDER,
    MANDATORY_HUMAN_REVIEW_SECTIONS,
    resolve_legacy,
    to_thai_numeral,
    validate_taxonomy_edit,
    TaxonomyEditError,
)
from app.export.docx_generator import TORContent
from app.export.export_plan import build_render_plan
from app.export.format_config import (
    BODY_FONT_SIZE_PT,
    HEADING_FONT_SIZE_PT,
    LINE_SPACING,
    PAGE_MARGIN_LEFT_RIGHT_CM,
    PAGE_MARGIN_TOP_BOTTOM_CM,
    apply_format,
)
from app.export.quality_gate import (
    GateWeakenedError,
    assert_gates_not_weakened,
    check_export_gates,
)
from app.export.render_plan import AppendixItem, NumberingScheme, coerce_numbering_scheme

_MAIN_KEYS = (
    "s1",
    "s2",
    "s3",
    "s4",
    "s5",
    "s6",
    "s7",
    "s8",
    "s9",
    "s10",
    "s11",
    "s13",
    "s15",
    "s16",
    "s17",
)
_TYPES = (
    "hire_develop",
    "hire_maintain",
    "lease_service",
    "buy_goods",
    "construction",
    "hire_consult",
    "hire_service",
)
_THAI_DIGITS = "๐๑๒๓๔๕๖๗๘๙"


def _to_int(num: str) -> int:
    mapped = "".join(str(_THAI_DIGITS.index(ch)) if ch in _THAI_DIGITS else ch for ch in num)
    return int(mapped.split(".", 1)[0])


@st.composite
def _draft_content(draw):
    project_type = draw(st.sampled_from(_TYPES))
    filled = draw(st.lists(st.sampled_from(_MAIN_KEYS), min_size=0, max_size=8, unique=True))
    sections = {key: "เนื้อหาหมวด " + key + "ก" * 8 for key in filled}
    subs: dict[str, dict[str, str]] = {}
    if draw(st.booleans()):
        subs["s4"] = {"s4.1": "ระบบงานปัจจุบันที่ต้อง archive", "s4.2": "เนื้อหา archive"}
        if draw(st.booleans()):
            subs["s4"]["functional"] = "งานพัฒนาระบบลงทะเบียน"
    if draw(st.booleans()):
        sections["s7"] = "สถานที่ติดตั้งอุปกรณ์"
    scheme = draw(st.sampled_from(list(NumberingScheme)))
    use_thai = draw(st.booleans())
    n_app = draw(st.integers(min_value=0, max_value=4))
    appendices = [
        AppendixItem(title=f"เอกสาร {index}", content=f"เนื้อหาภาคผนวก {index}")
        for index in range(1, n_app + 1)
    ]
    return TORContent(
        project_name="โครงการทดสอบ",
        project_type=project_type,
        sections=sections,
        sub_sections=subs,
        numbering_scheme=scheme,
        use_thai_numerals=use_thai,
        appendices=appendices,
    )


def test_format_config_matches_locked_layout():
    assert BODY_FONT_SIZE_PT == 16
    assert HEADING_FONT_SIZE_PT == 18
    assert LINE_SPACING == 1.0
    assert PAGE_MARGIN_TOP_BOTTOM_CM == 2.54
    assert PAGE_MARGIN_LEFT_RIGHT_CM == 1.91


def test_apply_format_skips_failed_fields():
    seen: list[str] = []

    def ok() -> None:
        seen.append("ok")

    def boom() -> None:
        raise RuntimeError("font missing")

    apply_format({"ok": ok, "boom": boom})
    assert seen == ["ok"]


def test_export_request_defaults_to_none_numbering():
    from app.schemas.export import ExportRequest

    assert ExportRequest().numbering_scheme == "none"


def test_thai_only_helpers_fail_closed():
    from app.services.thai_draft import (
        THAI_ONLY_RULES,
        ThaiOnlyNotAttachedError,
        attach_thai_only,
        detect_unauthorized_english,
        ensure_thai_only_attached,
    )

    assert "Price Only" in detect_unauthorized_english("ใช้วิธี Price Only ในการพิจารณา")
    assert not detect_unauthorized_english("ใช้เกณฑ์ราคาในระบบ e-GP ตาม PDPA")
    from app.services.thai_draft import sanitize_unauthorized_english

    cleaned = sanitize_unauthorized_english(
        "การจ่ายเงินแบ่งเป็นสามงวด ตาม Deliverables ที่ส่งมอบในระบบ e-GP "
        "เมื่อคณะกรรมการตรวจรับพัสดุได้ตรวจรับเรียบร้อยแล้วและเชื่อมโยงกับผลงานส่งมอบครบถ้วน"
        "ตามสัญญาจ้างพัฒนาระบบสารสนเทศเพื่อการบริหาร"
    )
    assert "Deliverables" not in cleaned
    assert "ผลงานส่งมอบ" in cleaned
    assert "e-GP" in cleaned
    assert not detect_unauthorized_english(cleaned)
    with pytest.raises(ThaiOnlyNotAttachedError):
        ensure_thai_only_attached("เขียนเป็นภาษาไทยราชการ")
    assert THAI_ONLY_RULES in attach_thai_only("เขียนเป็นภาษาไทยราชการ")


def test_coerce_numbering_scheme_invalid_falls_back():
    assert coerce_numbering_scheme("none") is NumberingScheme.NONE
    assert coerce_numbering_scheme("numbered_consecutive") is NumberingScheme.NUMBERED_CONSECUTIVE
    assert coerce_numbering_scheme("nope") is NumberingScheme.NUMBERED_CONSECUTIVE
    assert coerce_numbering_scheme(None) is NumberingScheme.NUMBERED_CONSECUTIVE


def test_appendix_divider_when_empty():
    from app.export.docx_generator import DOCXGenerator
    from app.export.pdf_generator import PDFGenerator
    import io
    from docx import Document

    content = TORContent(project_name="ทดสอบ", sections={"s1": "เนื้อหาความเป็นมา"})
    doc = Document(io.BytesIO(DOCXGenerator().generate(content)))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "ภาคผนวก" in text
    html = PDFGenerator()._build_html(content)
    assert "ภาคผนวก" in html


def test_resolve_legacy_archive_examples():
    assert resolve_legacy("s4", "s4.1", "hire_develop") is None
    assert resolve_legacy("s4", "s4.2", "buy_goods") is None
    mapped = resolve_legacy("s7", None, "hire_develop")
    assert mapped is not None
    assert mapped[0] == "scope"


def test_validate_taxonomy_edit_rejects_missing_core():
    with pytest.raises(TaxonomyEditError):
        validate_taxonomy_edit(core=list(CORE_SECTION_ORDER[1:]))


def test_generators_do_not_import_tor_sections():
    from app.export import docx_generator, pdf_generator, export_plan

    for module in (docx_generator, pdf_generator, export_plan):
        source = inspect.getsource(module)
        assert "tor_sections" not in source


def _padded_content() -> TORContent:
    sections = {
        "s1": "ก" * 420,
        "s2": "ข" * 160,
        "s3": "ค" * 420,
        "s4": "ง" * 620,
        "s5": "จ" * 130,
        "s6": "ฉ" * 80,
        "s8": "ช" * 160,
        "s10": "ซ" * 130,
        "s9": "ฌ" * 110,
        "s11": "ญ" * 90,
        "s17": "ด" * 70,
    }
    return TORContent(project_name="โครงการ", project_type="buy_goods", sections=sections)


def test_quality_gate_does_not_block_empty_appendices():
    content = _padded_content()
    approvals = {key: True for key in MANDATORY_HUMAN_REVIEW_SECTIONS}
    approvals.update({"s3": True, "s6": True, "s8": True, "s10": True, "s11": True})
    result = check_export_gates(content, approvals)
    assert result.ok
    assert not any("ภาคผนวก" in item for item in result.errors)


def test_quality_gate_blocks_short_section():
    content = _padded_content()
    content.sections["s1"] = "สั้น"
    approvals = {key: True for key in ("qualification", "budget", "payment", "penalty", "evaluation", "s3", "s6", "s8", "s10", "s11")}
    result = check_export_gates(content, approvals)
    assert not result.ok
    joined = " ".join(result.errors)
    assert "background" in joined
    assert "400" in joined


@pytest.mark.property
class TestTorOutputStandardization:
    @given(content=_draft_content())
    @settings(max_examples=100, deadline=None)
    def test_property_1_body_numbers_are_contiguous(self, content: TORContent):
        """Feature: tor-output-standardization, Property 1: body numbers 1..N contiguous"""
        content.numbering_scheme = NumberingScheme.NUMBERED_CONSECUTIVE
        plan = build_render_plan(content)
        numbers = [_to_int(section.number) for section in plan.body]
        assert numbers == list(range(1, len(numbers) + 1))

    @given(content=_draft_content())
    @settings(max_examples=100, deadline=None)
    def test_property_2_subsection_numbers_restart(self, content: TORContent):
        """Feature: tor-output-standardization, Property 2: subsection parent.child restart"""
        content.numbering_scheme = NumberingScheme.NUMBERED_CONSECUTIVE
        plan = build_render_plan(content)
        for section in plan.body:
            for index, sub in enumerate(section.subsections, start=1):
                assert sub.number.endswith(f".{to_thai_numeral(index) if content.use_thai_numerals else index}")
                assert sub.number.startswith(section.number)

    @given(content=_draft_content())
    @settings(max_examples=100, deadline=None)
    def test_property_3_empty_sections_skip_numbers(self, content: TORContent):
        """Feature: tor-output-standardization, Property 3: hidden rows do not consume numbers"""
        content.numbering_scheme = NumberingScheme.NUMBERED_CONSECUTIVE
        plan = build_render_plan(content)
        displayed = {section.storage_key for section in plan.body}
        for key, text in content.sections.items():
            if not str(text).strip() and key not in (content.sub_sections or {}):
                assert key not in displayed
        numbers = [_to_int(section.number) for section in plan.body]
        assert numbers == list(range(1, len(numbers) + 1))

    @given(content=_draft_content())
    @settings(max_examples=100, deadline=None)
    def test_property_4_none_keeps_order(self, content: TORContent):
        """Feature: tor-output-standardization, Property 4: none mode keeps order, empty numbers"""
        numbered = TORContent(
            project_name=content.project_name,
            project_type=content.project_type,
            sections=dict(content.sections),
            sub_sections=dict(content.sub_sections),
            numbering_scheme=NumberingScheme.NUMBERED_CONSECUTIVE,
            use_thai_numerals=content.use_thai_numerals,
            appendices=list(content.appendices),
        )
        none = TORContent(
            project_name=content.project_name,
            project_type=content.project_type,
            sections=dict(content.sections),
            sub_sections=dict(content.sub_sections),
            numbering_scheme=NumberingScheme.NONE,
            use_thai_numerals=content.use_thai_numerals,
            appendices=list(content.appendices),
        )
        numbered_plan = build_render_plan(numbered)
        none_plan = build_render_plan(none)
        assert [row.storage_key for row in numbered_plan.body] == [
            row.storage_key for row in none_plan.body
        ]
        assert all(row.number == "" for row in none_plan.body)
        for row in none_plan.body:
            assert all(sub.number == "" for sub in row.subsections)

    @given(st.text(min_size=1, max_size=24).filter(lambda value: value not in {"none", "numbered_consecutive"}))
    @settings(max_examples=100, deadline=None)
    def test_property_5_invalid_scheme_defaults(self, value: str):
        """Feature: tor-output-standardization, Property 5: invalid numbering → consecutive"""
        assert coerce_numbering_scheme(value) is NumberingScheme.NUMBERED_CONSECUTIVE

    @given(content=_draft_content())
    @settings(max_examples=100, deadline=None)
    def test_property_6_appendix_numbers_independent(self, content: TORContent):
        """Feature: tor-output-standardization, Property 6: appendix labels start at 1"""
        plan = build_render_plan(content)
        if not plan.appendices:
            return
        for index, item in enumerate(plan.appendices, start=1):
            expected = to_thai_numeral(index) if content.use_thai_numerals else str(index)
            assert item.number == f"ภาคผนวก {expected}"

    @given(content=_draft_content())
    @settings(max_examples=100, deadline=None)
    def test_property_7_body_numbers_ignore_appendix_count(self, content: TORContent):
        """Feature: tor-output-standardization, Property 7: body numbers independent of appendices"""
        content.numbering_scheme = NumberingScheme.NUMBERED_CONSECUTIVE
        empty = TORContent(
            project_name=content.project_name,
            project_type=content.project_type,
            sections=dict(content.sections),
            sub_sections=dict(content.sub_sections),
            numbering_scheme=NumberingScheme.NUMBERED_CONSECUTIVE,
            use_thai_numerals=content.use_thai_numerals,
            appendices=[],
        )
        with_apps = TORContent(
            project_name=content.project_name,
            project_type=content.project_type,
            sections=dict(content.sections),
            sub_sections=dict(content.sub_sections),
            numbering_scheme=NumberingScheme.NUMBERED_CONSECUTIVE,
            use_thai_numerals=content.use_thai_numerals,
            appendices=list(content.appendices)
            + [AppendixItem(title="เพิ่ม", content="รายละเอียด")],
        )
        left = [row.number for row in build_render_plan(empty).body]
        right = [row.number for row in build_render_plan(with_apps).body]
        assert left == right

    @given(st.booleans())
    @settings(max_examples=100, deadline=None)
    def test_property_8_zero_appendices_complete(self, extra_hitl: bool):
        """Feature: tor-output-standardization, Property 8: zero appendices stay complete"""
        content = _padded_content()
        approvals = {
            "qualification": True,
            "budget": True,
            "payment": True,
            "penalty": True,
            "evaluation": True,
            "s3": True,
            "s6": True,
            "s8": True,
            "s10": True,
            "s11": True,
        }
        if extra_hitl:
            approvals["ip_ownership"] = True
        result = check_export_gates(content, approvals)
        assert result.ok
        assert content.appendices == []

    @given(content=_draft_content())
    @settings(max_examples=100, deadline=None)
    def test_property_9_none_legacy_is_archived(self, content: TORContent):
        """Feature: tor-output-standardization, Property 9: resolve_legacy None is archived"""
        content.sub_sections = {
            "s4": {"s4.1": "เนื้อหาเก่า ๔.๑", "s4.2": "เนื้อหาเก่า ๔.๒"}
        }
        plan = build_render_plan(content)
        archived_keys = {(row.legacy_key, row.sub_key) for row in plan.archived}
        assert ("s4", "s4.1") in archived_keys
        assert ("s4", "s4.2") in archived_keys
        joined = " ".join(section.content for section in plan.body)
        joined += " ".join(sub.content for section in plan.body for sub in section.subsections)
        assert "เนื้อหาเก่า ๔.๑" not in joined
        assert "เนื้อหาเก่า ๔.๒" not in joined

    @given(st.sampled_from(_TYPES))
    @settings(max_examples=100, deadline=None)
    def test_property_10_legacy_with_target_is_kept(self, project_type: str):
        """Feature: tor-output-standardization, Property 10: mapped legacy is preserved"""
        content = TORContent(
            project_name="โครงการ",
            project_type=project_type,
            sections={"s12": "บัญชีเอกสารที่ต้องยื่นประกอบการเสนอราคา"},
            numbering_scheme=NumberingScheme.NONE,
        )
        plan = build_render_plan(content)
        joined = " ".join(
            [section.content for section in plan.body]
            + [sub.content for section in plan.body for sub in section.subsections]
        )
        assert "บัญชีเอกสารที่ต้องยื่น" in joined
        assert not any(row.legacy_key == "s12" for row in plan.archived)

    @given(content=_draft_content())
    @settings(max_examples=100, deadline=None)
    def test_property_11_docx_pdf_share_plan(self, content: TORContent):
        """Feature: tor-output-standardization, Property 11: DOCX and PDF share RenderPlan"""
        from app.export.docx_generator import DOCXGenerator
        from app.export.pdf_generator import PDFGenerator

        plan = build_render_plan(content)
        html = PDFGenerator()._build_html(content)
        for section in plan.body:
            heading = f"{section.number}. {section.label}" if section.number else section.label
            assert heading in html
        assert "ภาคผนวก" in html
        import io
        from docx import Document

        doc = Document(io.BytesIO(DOCXGenerator().generate(content)))
        text = "\n".join(p.text for p in doc.paragraphs)
        for section in plan.body:
            heading = f"{section.number}. {section.label}" if section.number else section.label
            assert heading in text
        assert "ภาคผนวก" in text

    @given(
        st.integers(min_value=1, max_value=50),
        st.sampled_from(sorted(CRITICAL_SECTIONS_MIN_LENGTH)),
    )
    @settings(max_examples=100, deadline=None)
    def test_property_12_gates_not_weakened(self, delta: int, key: str):
        """Feature: tor-output-standardization, Property 12: gates cannot weaken"""
        after = dict(CRITICAL_SECTIONS_MIN_LENGTH)
        after[key] = max(0, after[key] - delta)
        if after[key] < CRITICAL_SECTIONS_MIN_LENGTH[key]:
            with pytest.raises(GateWeakenedError):
                assert_gates_not_weakened(CRITICAL_SECTIONS_MIN_LENGTH, after)
        else:
            assert_gates_not_weakened(CRITICAL_SECTIONS_MIN_LENGTH, after)
        weakened_hitl = set(MANDATORY_HUMAN_REVIEW_SECTIONS) - {"budget"}
        with pytest.raises(GateWeakenedError):
            assert_gates_not_weakened(
                CRITICAL_SECTIONS_MIN_LENGTH,
                after_hitl=weakened_hitl,
            )

    @given(st.sampled_from(sorted(CRITICAL_SECTIONS_MIN_LENGTH)), st.integers(min_value=0, max_value=20))
    @settings(max_examples=100, deadline=None)
    def test_property_13_short_critical_blocks_export(self, semantic: str, length: int):
        """Feature: tor-output-standardization, Property 13: short critical section blocks export"""
        from app.domain.section_profile import SEMANTIC_TO_STORAGE, category_for_project
        from app.domain.tor_taxonomy import section_order

        content = _padded_content()
        if semantic not in section_order(category_for_project(content.project_type)):
            return
        storage = SEMANTIC_TO_STORAGE.get(semantic, semantic)
        content.sections[storage] = "ก" * length
        if semantic == "scope":
            content.sub_sections = {}
        minimum = CRITICAL_SECTIONS_MIN_LENGTH[semantic]
        if length >= minimum:
            return
        approvals = {key: True for key in ("qualification", "budget", "payment", "penalty", "evaluation", "s3", "s6", "s8", "s10", "s11")}
        result = check_export_gates(content, approvals)
        assert not result.ok
        joined = " ".join(result.errors)
        assert semantic in joined
        assert str(minimum) in joined
        assert str(length) in joined
