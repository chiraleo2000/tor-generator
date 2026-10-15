"""Table and image layout extraction.

Vector PDFs keep table grids as markdown in reading order. Scanned pages OCR
after table rules are removed, then rebuild the same grid from word boxes.
"""

from pathlib import Path

import fitz
import pytest
from docx import Document
from docx.shared import Inches

from app.rag.extraction import (
    OcrWord,
    _tesseract_executable,
    extract_docx,
    extract_pdf,
    words_to_layout_text,
)

THAI_FONT = Path(r"C:\Windows\Fonts\tahoma.ttf")


def _word(text: str, left: int, top: int, width: int = 24, height: int = 16) -> OcrWord:
    return OcrWord(text=text, left=left, top=top, width=width, height=height)


def test_word_boxes_become_a_markdown_table_between_prose():
    words = [
        _word("ขอบเขต", 10, 10, 48),
        _word("รายการ", 10, 50, 56),
        _word("จำนวน", 220, 50, 48),
        _word("หน่วย", 420, 50, 40),
        _word("คอมพิวเตอร์", 10, 90, 96),
        _word("10", 220, 90, 24),
        _word("เครื่อง", 420, 90, 48),
        _word("หลังตาราง", 10, 140, 72),
    ]
    text = words_to_layout_text(words)
    assert "| รายการ | จำนวน | หน่วย |" in text
    assert "| คอมพิวเตอร์ | 10 | เครื่อง |" in text
    assert text.index("ขอบเขต") < text.index("| รายการ |")
    assert text.index("เครื่อง |") < text.index("หลังตาราง")


def test_thai_glyphs_on_one_line_stay_one_phrase():
    pieces = ["ขอ", "บ", "เข", "ต", "ของ", "งาน"]
    words = [_word(piece, 20 + index * 22, 30, 18) for index, piece in enumerate(pieces)]
    assert words_to_layout_text(words) == "ขอบเขตของงาน"


def test_nikhahit_and_sara_aa_join_into_sara_am():
    words = [
        _word("จ", 10, 10, 14),
        _word("\u0e4d", 24, 8, 10),
        _word("\u0e32", 34, 10, 14),
        _word("น", 50, 10, 14),
        _word("ว", 66, 10, 14),
        _word("น", 82, 10, 14),
    ]
    assert words_to_layout_text(words) == "จำนวน"


def _draw_thai_table(page: fitz.Page) -> None:
    page.insert_font(fontname="thai", fontfile=str(THAI_FONT))
    page.insert_text((72, 56), "ขอบเขตของงาน", fontname="thai", fontsize=16)
    rows = (["รายการ", "จำนวน", "หน่วย"], ["คอมพิวเตอร์", "10", "เครื่อง"])
    origin_x, origin_y, col_w, row_h = 72, 80, 150, 36
    for row_index, row in enumerate(rows):
        for col_index, cell in enumerate(row):
            rect = fitz.Rect(
                origin_x + col_index * col_w,
                origin_y + row_index * row_h,
                origin_x + (col_index + 1) * col_w,
                origin_y + (row_index + 1) * row_h,
            )
            page.draw_rect(rect)
            page.insert_textbox(
                rect + (6, 8, -4, -4),
                cell,
                fontname="thai",
                fontsize=12,
            )
    page.insert_text((72, 180), "หลังตาราง", fontname="thai", fontsize=14)


@pytest.mark.skipif(not THAI_FONT.is_file(), reason="Thai font is not installed")
def test_vector_pdf_table_is_markdown_once_and_in_order(tmp_path: Path):
    doc = fitz.open()
    _draw_thai_table(doc.new_page())
    path = tmp_path / "table.pdf"
    doc.save(path)
    doc.close()

    text = extract_pdf(str(path)).text
    assert "| รายการ | จำนวน | หน่วย |" in text
    assert "| คอมพิวเตอร์ | 10 | เครื่อง |" in text
    assert text.count("คอมพิวเตอร์") == 1
    assert text.index("ขอบเขตของงาน") < text.index("| รายการ |")
    assert text.index("| เครื่อง |") < text.index("หลังตาราง")


def _tesseract_installed() -> bool:
    exe = Path(_tesseract_executable())
    return exe.is_file()


@pytest.mark.skipif(not THAI_FONT.is_file(), reason="Thai font is not installed")
@pytest.mark.skipif(not _tesseract_installed(), reason="Tesseract is not installed")
def test_scanned_pdf_table_ocr_keeps_cells_and_order(tmp_path: Path):
    source = fitz.open()
    page = source.new_page()
    _draw_thai_table(page)
    png = tmp_path / "scan.png"
    page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).save(png)
    source.close()

    scanned = fitz.open()
    image_page = scanned.new_page()
    image_page.insert_image(fitz.Rect(36, 36, 560, 360), filename=str(png))
    path = tmp_path / "scanned.pdf"
    scanned.save(path)
    scanned.close()

    text = extract_pdf(str(path)).text
    tokens = (
        "ขอบเขตของงาน",
        "รายการ",
        "จำนวน",
        "หน่วย",
        "คอมพิวเตอร์",
        "10",
        "เครื่อง",
        "หลังตาราง",
    )
    for token in tokens:
        assert token in text
    assert text.index("ขอบเขตของงาน") < text.index("รายการ") < text.index("คอมพิวเตอร์")
    assert text.index("คอมพิวเตอร์") < text.index("หลังตาราง")
    assert "| รายการ | จำนวน | หน่วย |" in text
    assert "| คอมพิวเตอร์ | 10 | เครื่อง |" in text


def test_docx_merged_cells_stay_in_column_order(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("ก่อนตาราง")
    table = doc.add_table(rows=2, cols=3)
    table.rows[0].cells[0].text = "หัวข้อ"
    table.rows[0].cells[1].text = "จำนวน"
    table.rows[0].cells[2].text = "หน่วย"
    merged = table.rows[1].cells[0]
    merged.merge(table.rows[1].cells[1])
    merged.text = "รวม"
    table.rows[1].cells[2].text = "ชุด"
    doc.add_paragraph("หลังตาราง")
    path = tmp_path / "merged.docx"
    doc.save(path)

    text = extract_docx(str(path)).text
    assert "| หัวข้อ | จำนวน | หน่วย |" in text
    assert text.count("รวม") == 1
    assert "ชุด" in text
    assert text.index("ก่อนตาราง") < text.index("หัวข้อ") < text.index("หลังตาราง")


def test_docx_picture_text_is_inserted_beside_the_paragraph(tmp_path: Path, monkeypatch):
    blank = fitz.open()
    page = blank.new_page(width=120, height=60)
    page.draw_rect(page.rect, color=(1, 1, 1), fill=(1, 1, 1))
    png = tmp_path / "figure.png"
    page.get_pixmap(alpha=False).save(png)
    blank.close()
    monkeypatch.setattr(
        "app.rag.extraction._ocr_image_bytes",
        lambda _blob: ("ข้อความในภาพ", None),
    )
    doc = Document()
    doc.add_paragraph("ก่อนภาพ")
    doc.add_picture(str(png), width=Inches(2))
    doc.add_paragraph("หลังภาพ")
    path = tmp_path / "picture.docx"
    doc.save(path)

    text = extract_docx(str(path)).text
    assert text.index("ก่อนภาพ") < text.index("ข้อความในภาพ") < text.index("หลังภาพ")
